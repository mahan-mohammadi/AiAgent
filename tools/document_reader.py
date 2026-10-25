

from pathlib import Path

from docx import Document
from langchain_core.tools import tool
from pypdf import PdfReader


DOCUMENTS_DIR = (
    Path(__file__).resolve().parent.parent / "documents"
).resolve()

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
MAX_OUTPUT_CHARS = 12000  # what read_document returns to the model


def resolve_document(filename: str) -> Path:
    """Resolve a document without allowing access outside documents/."""
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)

    relative_path = Path(filename)

    if relative_path.is_absolute():
        raise ValueError("Use a relative filename inside documents/.")

    path = (DOCUMENTS_DIR / relative_path).resolve()

    if not path.is_relative_to(DOCUMENTS_DIR):
        raise ValueError("Access outside documents/ is not allowed.")

    return path


def load_text(path: Path) -> str:
    """Extract the full plain text of a supported document.

    No size limit is applied here. Callers decide what to do with it:
    read_document truncates it, a RAG indexer would split it into chunks.
    """
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")

    if suffix == ".docx":
        doc = Document(path)
        content = "\n".join(paragraph.text for paragraph in doc.paragraphs)

        for table in doc.tables:
            for row in table.rows:
                content += "\n" + " | ".join(cell.text for cell in row.cells)

        return content

    if suffix == ".pdf":
        pdf = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in pdf.pages)

    raise ValueError(f"Unsupported file type: {suffix or '(none)'}")


@tool
def read_document(filename: str) -> str:
    """Read a text, Markdown, PDF, or Word document from documents/.
    Provide a relative filename, such as 'notes.md' or 'reports/paper.pdf'.
    Use this to answer questions about existing documents.
    """
    try:
        path = resolve_document(filename)

        if not path.is_file():
            return f"Error: document not found: {filename}"

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            return "Error: supported formats are .txt, .md, .pdf, and .docx."

        if path.stat().st_size > MAX_FILE_SIZE:
            return "Error: document exceeds the 10 MB limit."

        content = load_text(path)

        if not content.strip():
            return "The document contains no extractable text."

        # Keep tool output bounded for the model's context window.
        if len(content) > MAX_OUTPUT_CHARS:
            content = content[:MAX_OUTPUT_CHARS] + "\n\n[Content truncated.]"

        return f"Document: {path.name}\n\n{content}"

    except Exception as exc:
        return f"Document read failed: {type(exc).__name__}: {exc}"
